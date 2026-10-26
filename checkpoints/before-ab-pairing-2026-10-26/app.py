from __future__ import annotations

import base64
import csv
import json
import re
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
ROSTER = [
    {"name": "魏沐翔", "attending": True, "noncritical": True, "supervisor": False},
    {"name": "張德忠", "attending": True, "noncritical": False, "supervisor": True},
    {"name": "吳貴立", "attending": True, "noncritical": False, "supervisor": True},
    {"name": "羅俊麟", "attending": True, "noncritical": False, "supervisor": True},
    {"name": "黃偉均", "attending": True, "noncritical": False, "supervisor": True},
    {"name": "吳東穎", "attending": True, "noncritical": False, "supervisor": True},
    {"name": "楊捷先", "attending": True, "noncritical": False, "supervisor": True},
    {"name": "陳宗偉", "attending": True, "noncritical": False, "supervisor": True},
    {"name": "余中仁", "attending": True, "noncritical": False, "supervisor": True},
]
HEADERS = {
    "date": "預定施工日期",
    "department": "部門別",
    "job_no": "施工號碼",
    "contractor": "施工者",
    "contract": "契約名稱",
    "crew": "工作班",
    "time": "施工時間",
    "description": "工作內容",
    "location": "施工地點",
}
GRID_FULL = re.compile(
    r"(?<![A-Z0-9])[A-Z]\s*\d{4}\s*[A-Z]{2}\s*\d{2}(?:\s*\d{2})?(?![A-Z0-9])",
    re.IGNORECASE,
)
GRID_TOKEN = re.compile(
    r"(?<![A-Z0-9])(?:[A-Z]\s*\d{4}\s*[A-Z]{2}\s*\d{2}(?:\s*\d{2})?|[A-Z]{2}\s*\d{2}(?:\s*\d{2})?)(?![A-Z0-9])",
    re.IGNORECASE,
)

# Taipower's grid IDs encode TWD67 TM2 meter positions. These sector origins
# follow the published open conversion algorithm; the last two digits identify
# a 10 m cell and the optional final two digits identify a 1 m cell.
GRID_BASELINES = {
    "S": (10000, 2894000), "Y": (275000, 2564000),
    "X": (275000, 2614000), "Z": (10000, 2675800),
}
TAIWAN_SECTORS = ("_ABC", "_DEF", "_GH_", "JKL_", "MNO_", "PQR_", "_TU_", "_VW")
_bottom = 2_800_000
for _row_index, _row in enumerate(TAIWAN_SECTORS):
    _south = _bottom - (_row_index + 1) * 50_000
    for _col_index, _letter in enumerate(_row):
        if _letter.isalpha():
            GRID_BASELINES[_letter] = (10_000 + (_col_index + 1) * 80_000, _south)


def decode_grid_code(code):
    """Return the center of a Taipower grid cell in TWD67 TM2 meters."""
    match = re.fullmatch(r"([A-HJ-Z])(\d{2})(\d{2})([A-H])([A-E])(\d)(\d)(?:(\d)(\d))?", code)
    if not match:
        return None
    sector, east_map, north_map, east_letter, north_letter, east_10, north_10, east_1, north_1 = match.groups()
    baseline = GRID_BASELINES.get(sector)
    if baseline is None:
        return None
    if sector == "Z":
        east_map = str(50 + (int(east_map) + 50) % 100)
    east = baseline[0] + 800 * int(east_map) + 100 * (ord(east_letter) - ord("A")) + 10 * int(east_10)
    north = baseline[1] + 500 * int(north_map) + 100 * (ord(north_letter) - ord("A")) + 10 * int(north_10)
    precision = 1 if east_1 is not None else 10
    if east_1 is not None:
        east += int(east_1)
        north += int(north_1)
    return {"x": east + precision / 2, "y": north + precision / 2, "crs": "TWD67 TM2", "precision": precision}


def normalize_grid_code(value):
    """Normalize whitespace (including zero-width spaces) inside Taipower IDs."""
    return re.sub(r"[\s\u200b\ufeff]+", "", str(value or "")).upper()


def add_grid_coordinates(jobs):
    """Decode all grid IDs; use the local CSV only to enrich pole metadata."""
    source = ROOT / "all.csv"
    wanted = {point["code"] for job in jobs for point in job["gridCodes"]}
    for job in jobs:
        for grid_code in job["gridCodes"]:
            decoded = decode_grid_code(grid_code["code"])
            if decoded:
                grid_code["gridCoordinate"] = decoded
    if not wanted or not source.is_file():
        return {"matched": 0, "decoded": sum(bool(decode_grid_code(code)) for code in wanted), "total": len(wanted), "available": source.is_file()}

    points = {}
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            code = normalize_grid_code(row.get("圖號座標", ""))
            if code not in wanted or code in points:
                continue
            try:
                x, y = float(row["TWD_97_X"]), float(row["TWD_97_Y"])
            except (KeyError, TypeError, ValueError):
                continue
            points[code] = {
                "x": x,
                "y": y,
                "county": row.get("縣市", ""),
                "district": row.get("行政區", ""),
                "poleNo": row.get("桿號", ""),
            }
            if len(points) == len(wanted):
                break

    for job in jobs:
        for grid_code in job["gridCodes"]:
            coordinate = points.get(grid_code["code"])
            if coordinate:
                grid_code["coordinate"] = coordinate
    return {"matched": len(points), "decoded": sum(bool(decode_grid_code(code)) for code in wanted), "total": len(wanted), "available": True}


def extract_grid_codes(work_content, location):
    """Collect full Taipower grid IDs and expand linked short suffixes in location text."""
    found = []
    seen = set()

    def add(code, expanded=False):
        code = normalize_grid_code(code)
        if code not in seen:
            seen.add(code)
            found.append({"code": code, "expanded": expanded, "representative": not found})

    work_text = str(work_content or "")
    for match in GRID_FULL.finditer(work_text):
        add(match.group())

    location_text = str(location or "")
    last_grid_prefix = None
    for match in GRID_TOKEN.finditer(location_text):
        token = normalize_grid_code(match.group())
        if re.fullmatch(r"[A-Z]\d{4}[A-Z]{2}\d{2}(?:\d{2})?", token):
            add(token)
            last_grid_prefix = token[:5]
        elif last_grid_prefix:
            expanded = last_grid_prefix + token
            if re.fullmatch(r"[A-Z]\d{4}[A-Z]{2}\d{2}(?:\d{2})?", expanded):
                add(expanded, expanded=True)
    return found


def date_key(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip().splitlines()[0]
    match = re.match(r"^(\d{2,4})[/-](\d{1,2})[/-](\d{1,2})$", text)
    if not match:
        return text
    year, month, day = map(int, match.groups())
    if year < 1911:
        year += 1911
    return f"{year:04d}-{month:02d}-{day:02d}"


def read_excel(payload):
    data = base64.b64decode(payload)
    workbook = load_workbook(filename=__import__("io").BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    header_row = None
    headers = []
    for row_num, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 15), values_only=True), start=1):
        values = [str(value).strip() if value is not None else "" for value in row]
        if HEADERS["date"] in values and HEADERS["crew"] in values:
            header_row, headers = row_num, values
            break
    if header_row is None:
        raise ValueError("找不到「預定施工日期」與「工作班」欄位，請確認上傳的是台電施工資料匯出檔。")
    indexes = {key: headers.index(label) for key, label in HEADERS.items() if label in headers}
    jobs = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        get = lambda key: row[indexes[key]] if key in indexes and indexes[key] < len(row) else None
        work_date = date_key(get("date"))
        crew = str(get("crew") or "").strip()
        if not work_date or not crew:
            continue
        description = str(get("description") or "").strip()
        location = str(get("location") or "").strip()
        work_summary = next((line.strip() for line in description.splitlines() if line.strip()), "")
        jobs.append({
            "date": work_date,
            "department": str(get("department") or "").replace("\n", " ").strip(),
            "jobNo": str(get("job_no") or "").strip(),
            "contractor": str(get("contractor") or "").strip(),
            "contract": str(get("contract") or "").strip(),
            "crew": crew.replace("\n", " ").strip(),
            "time": str(get("time") or "").replace("\n", "–").strip(),
            "description": description,
            "workSummary": work_summary,
            "location": location,
            "gridCodes": extract_grid_codes(description, location),
        })
    return jobs


class Handler(BaseHTTPRequestHandler):
    def _send(self, status, body, content_type="application/json; charset=utf-8"):
        encoded = body if isinstance(body, bytes) else body.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(encoded)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(encoded)

    def do_GET(self):
        path = urlparse(self.path).path
        if path in ("/", "/index.html"):
            return self._send(200, (ROOT / "index.html").read_bytes(), "text/html; charset=utf-8")
        if path == "/api/roster":
            return self._send(200, json.dumps(ROSTER, ensure_ascii=False))
        self._send(404, json.dumps({"error": "找不到頁面"}, ensure_ascii=False))

    def do_POST(self):
        if urlparse(self.path).path != "/api/parse":
            return self._send(404, json.dumps({"error": "找不到 API"}, ensure_ascii=False))
        try:
            length = int(self.headers.get("Content-Length", "0"))
            body = json.loads(self.rfile.read(length))
            jobs = read_excel(body.get("file", ""))
            coordinate_summary = add_grid_coordinates(jobs)
            self._send(200, json.dumps({"jobs": jobs, "coordinates": coordinate_summary}, ensure_ascii=False))
        except Exception as exc:
            self._send(400, json.dumps({"error": str(exc)}, ensure_ascii=False))

    def log_message(self, *_args):
        pass


if __name__ == "__main__":
    server = ThreadingHTTPServer(("127.0.0.1", 8000), Handler)
    print("自動派工雛形已啟動：http://127.0.0.1:8000")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n伺服器已停止")
